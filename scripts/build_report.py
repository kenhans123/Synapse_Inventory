#!/usr/bin/env python3
"""
Build the Synapse datamart object inventory workbook from inventory.json (see extract.py).

  python build_report.py --in inventory.json --out DM_Sales_Object_Inventory.xlsx [--target databricks|fabric]
                         [--no-definitions] [--client "CCEP"]
"""
import argparse
import datetime as dt
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _bootstrap  # noqa: E402,F401  (bundled openpyxl fallback - no pip install needed)

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from analyze import FEATURES, FEATURE_BY_KEY, RATING_BANDS, analyze  # noqa: E402
import lineage as lin  # noqa: E402

TYPE_ORDER = ["Stored Procedure", "Function", "Materialized View", "View", "Databricks Mirror", "External Table"]
SHEET_FOR = {"Stored Procedure": "Stored Procedures", "Function": "Functions",
             "Materialized View": "Materialized Views", "View": "Views",
             "Databricks Mirror": "Databricks Mirror", "External Table": "External Tables"}
RATINGS = ["Low", "Medium", "High", "Very High"]

FONT = "Arial"
NAVY, TEAL, LIGHT, GRID = "1F3864", "2E75B6", "EEF3FA", "BFC9D9"
RATING_FILL = {"Low": "C6EFCE", "Medium": "FFEB9C", "High": "F8CBAD", "Very High": "FFC7CE"}
RATING_FONT = {"Low": "006100", "Medium": "7F6000", "High": "843C0C", "Very High": "9C0006"}

thin = Side(style="thin", color=GRID)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HDR_FILL = PatternFill("solid", fgColor=NAVY)
BODY = Font(name=FONT, size=10)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def txt(v):
    if v is None:
        return None
    if isinstance(v, (list, tuple, set)):
        v = "\n".join(str(x) for x in v)
    if isinstance(v, bool):
        return "Y" if v else "N"
    if isinstance(v, (int, float)):
        return v
    v = ILLEGAL_CHARACTERS_RE.sub("", str(v))
    return v[:32000] + " …[truncated]" if len(v) > 32000 else v


def put(ws, row, col, value, font=BODY, align=None, fill=None, fmt=None, as_text=False):
    c = ws.cell(row=row, column=col)
    c.value = value
    if as_text and isinstance(value, str):
        c.data_type = "s"
    c.font = font
    if align:
        c.alignment = align
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    return c


def header_row(ws, row, headers, widths=None):
    for i, h in enumerate(headers, 1):
        c = put(ws, row, i, h, HDR_FONT, Alignment(wrap_text=True, vertical="center"), NAVY)
        c.border = BORDER
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def add_table(ws, name, first_row, last_row, ncols):
    if last_row <= first_row:
        return
    ref = f"A{first_row}:{get_column_letter(ncols)}{last_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


def rating_cf(ws, col_letter, first, last):
    if last < first:
        return
    rng = f"{col_letter}{first}:{col_letter}{last}"
    for r in RATINGS:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{r}"'],
                                                      fill=PatternFill("solid", fgColor=RATING_FILL[r]),
                                                      font=Font(name=FONT, color=RATING_FONT[r], bold=True)))


def title_block(ws, title, subtitle):
    put(ws, 1, 1, title, Font(name=FONT, bold=True, size=14, color=NAVY))
    put(ws, 2, 1, subtitle, Font(name=FONT, italic=True, size=9, color="595959"))


# --------------------------------------------------------------------------- sheets
def write_definitions(wb, objects):
    ws = wb.create_sheet("Definitions")
    ws.sheet_properties.tabColor = "7F7F7F"
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 150
    code_font = Font(name="Consolas", size=9)
    row = 1
    anchors = {}
    for o in objects:
        anchors[o["full_name"]] = row
        put(ws, row, 1, None, fill=NAVY)
        put(ws, row, 2, f"{o['full_name']}   ·   {o['object_type']}   ·   {o.get('lines', 0)} lines   ·   source: {o.get('source')}",
            Font(name=FONT, bold=True, color="FFFFFF", size=10), fill=NAVY)
        row += 1
        lines = (o.get("definition") or "-- definition not available").replace("\r\n", "\n").replace("\r", "\n").split("\n")
        while lines and not lines[-1].strip():
            lines.pop()
        for i, ln in enumerate(lines, 1):
            put(ws, row, 1, i, Font(name="Consolas", size=8, color="808080"))
            put(ws, row, 2, txt(ln.replace("\t", "    ")) or None, code_font, as_text=True)  # code starting '=' stays text
            row += 1
        row += 1
    ws.freeze_panes = "A1"
    return anchors


def link(ws, row, col, text, anchor_row):
    c = put(ws, row, col, text, Font(name=FONT, size=10, color="0563C1", underline="single"), TOP)
    if anchor_row:
        c.hyperlink = f"#Definitions!B{anchor_row}"
    return c


def write_inventory(wb, objects, anchors):
    ws = wb.create_sheet("Inventory", 1)
    ws.sheet_properties.tabColor = TEAL
    headers = ["#", "Full Name", "Schema", "Object Name", "Object Type", "Lines", "Complexity Score", "Rating",
               "Referenced By (count)", "Reads", "Writes", "Calls", "Key Features", "Migration Notes",
               "Created", "Modified", "Source"]
    widths = [5, 38, 12, 34, 18, 8, 11, 11, 12, 34, 30, 30, 32, 60, 17, 17, 28]
    header_row(ws, 1, headers, widths)
    for i, o in enumerate(objects, 2):
        vals = [i - 1, None, o["schema"], o["name"], o["object_type"], o.get("lines"), o.get("complexity_score"),
                o.get("rating"), f"=COUNTIF(Dependencies!$C:$C,B{i})", txt(o.get("reads")), txt(o.get("writes")),
                txt(o.get("calls")), ", ".join(FEATURE_BY_KEY[f][1] for f in o.get("features", [])),
                o.get("migration_notes"), o.get("created"), o.get("modified"), o.get("source")]
        for j, v in enumerate(vals, 1):
            if j == 2:
                link(ws, i, 2, o["full_name"], anchors.get(o["full_name"]))
                continue
            put(ws, i, j, txt(v) if j != 9 else v, BODY, WRAP_TOP if j in (10, 11, 12, 13, 14) else TOP,
                fmt="0.0" if j == 7 else None)
    last = len(objects) + 1
    add_table(ws, "tblInventory", 1, last, len(headers))
    rating_cf(ws, "H", 2, last)
    ws.freeze_panes = "C2"
    return ws


def write_type_sheet(wb, otype, objs, anchors):
    ws = wb.create_sheet(SHEET_FOR[otype])
    ws.sheet_properties.tabColor = {"Stored Procedure": "2E75B6", "Function": "548235", "Materialized View": "BF8F00",
                                    "View": "7030A0", "Databricks Mirror": "C55A11", "External Table": "C55A11"}[otype]
    base = [("#", 5), ("Full Name (→ definition)", 40), ("Schema", 12), ("Name", 34)]
    tail = [("Lines", 8), ("Complexity Score", 11), ("Rating", 11), ("Key Features", 32), ("Migration Notes", 60),
            ("Modified", 17)]
    if otype == "Stored Procedure":
        extra = [("Parameters", 34), ("Param Count", 9), ("Reads", 34), ("Writes", 32), ("Calls", 30)]
        getx = lambda o: [txt(o.get("parameters")), len(o.get("parameters") or []), txt(o.get("reads")),
                          txt(o.get("writes")), txt(o.get("calls"))]
    elif otype == "Function":
        extra = [("Function Kind", 18), ("Parameters", 30), ("Returns", 16), ("Reads", 34)]
        getx = lambda o: [o.get("function_kind"), txt(o.get("parameters")), o.get("returns"), txt(o.get("reads"))]
    elif otype == "Materialized View":
        extra = [("Distribution", 20), ("Base Objects", 36), ("GROUP BY", 9), ("Aggregates", 22)]
        getx = lambda o: [o.get("distribution"), txt(o.get("reads")), "Y" if o.get("has_group_by") else "N",
                          o.get("aggregates")]
    elif otype == "View":
        extra = [("Base Objects", 40), ("Base Object Count", 10), ("Has Aggregation", 10)]
        getx = lambda o: [txt(o.get("reads")), len(o.get("reads") or []), "Y" if o.get("has_group_by") else "N"]
    else:  # Databricks Mirror / External Table
        extra = [("Implemented As", 15), ("Location", 40), ("Data Source", 22), ("File Format", 18),
                 ("Format Type", 11), ("Columns", 9), ("Mirror Evidence", 40), ("Base Objects", 30)]
        getx = lambda o: [o.get("original_type") or o["object_type"], o.get("location"), o.get("data_source"),
                          o.get("file_format"), o.get("format_type"), o.get("column_count"),
                          o.get("mirror_evidence"), txt(o.get("reads"))]
    cols = base + extra + tail
    title_block(ws, f"{SHEET_FOR[otype]} ({len(objs)})", "Click a name to jump to its full definition. "
                "Rating colours: green Low · yellow Medium · orange High · red Very High.")
    header_row(ws, 4, [c[0] for c in cols], [c[1] for c in cols])
    for i, o in enumerate(objs, 5):
        vals = [i - 4, None, o["schema"], o["name"]] + getx(o) + [
            o.get("lines"), o.get("complexity_score"), o.get("rating"),
            ", ".join(FEATURE_BY_KEY[f][1] for f in o.get("features", [])), o.get("migration_notes"), o.get("modified")]
        for j, v in enumerate(vals, 1):
            if j == 2:
                link(ws, i, 2, o["full_name"], anchors.get(o["full_name"]))
            else:
                put(ws, i, j, txt(v), BODY, WRAP_TOP)
    last = len(objs) + 4
    if not objs:
        put(ws, 5, 1, f"No {SHEET_FOR[otype].lower()} found in this datamart.",
            Font(name=FONT, italic=True, size=10, color="595959"))
    tname = "tbl" + SHEET_FOR[otype].replace(" ", "")
    add_table(ws, tname, 4, last, len(cols))
    rating_col = get_column_letter(len(base) + len(extra) + 3)
    rating_cf(ws, rating_col, 5, last)
    ws.freeze_panes = "C5"


def write_features(wb, objects):
    ws = wb.create_sheet("Feature Matrix")
    ws.sheet_properties.tabColor = "8EA9DB"
    keys = [f[0] for f in FEATURES]
    headers = ["Full Name", "Object Type", "Rating"] + [FEATURE_BY_KEY[k][1] for k in keys]
    header_row(ws, 1, headers, [38, 18, 11] + [9] * len(keys))
    ws.row_dimensions[1].height = 90
    for c in range(4, len(headers) + 1):
        ws.cell(row=1, column=c).alignment = Alignment(text_rotation=90, horizontal="center", vertical="bottom", wrap_text=True)
    for i, o in enumerate(objects, 2):
        put(ws, i, 1, o["full_name"]); put(ws, i, 2, o["object_type"]); put(ws, i, 3, o["rating"])
        for j, k in enumerate(keys, 4):
            put(ws, i, j, "Y" if k in o.get("features", []) else None, Font(name=FONT, bold=True, size=10, color="1F3864"),
                Alignment(horizontal="center"))
    last = len(objects) + 1
    rng = f"D2:{get_column_letter(len(headers))}{max(last, 2)}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Y"'], fill=PatternFill("solid", fgColor="BDD7EE")))
    rating_cf(ws, "C", 2, last)
    add_table(ws, "tblFeatures", 1, last, len(headers))
    ws.freeze_panes = "D2"
    return keys


def write_dependencies(wb, deps):
    ws = wb.create_sheet("Dependencies")
    ws.sheet_properties.tabColor = "A5A5A5"
    headers = ["Object", "Object Type", "Referenced Object", "Referenced Type", "Operations", "Resolved In Datamart"]
    header_row(ws, 1, headers, [40, 18, 42, 22, 16, 12])
    for i, d in enumerate(deps, 2):
        for j, k in enumerate(["object", "object_type", "referenced", "referenced_type", "operations", "resolved"], 1):
            put(ws, i, j, txt(d[k]))
    add_table(ws, "tblDependencies", 1, len(deps) + 1, len(headers))
    ws.freeze_panes = "A2"


def write_summary(wb, meta, objects, feature_keys, target, client):
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    widths = [30, 12, 10, 10, 10, 10, 12, 14, 4, 34, 12, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    put(ws, 1, 1, f"{meta.get('datamart')} — Synapse Datamart Object Inventory", Font(name=FONT, bold=True, size=16, color=NAVY))
    src = (f"{meta.get('server')} / {meta.get('database')}" if meta.get("mode") == "live" else meta.get("path"))
    info = [("Client / engagement", client or "—"), ("Source", f"{meta.get('mode')} · {src}"),
            ("Extracted", meta.get("extracted_at")), ("Report generated", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("Migration target (notes)", "Azure Databricks (Unity Catalog / Delta)" if target == "databricks" else "Microsoft Fabric (Warehouse / Lakehouse)")]
    for r, (k, v) in enumerate(info, 3):
        put(ws, r, 1, k, Font(name=FONT, bold=True, size=10, color="595959"))
        put(ws, r, 2, v)

    # object counts x rating
    r0 = 10
    put(ws, r0 - 1, 1, "Objects by type and migration complexity", Font(name=FONT, bold=True, size=12, color=NAVY))
    hdr = ["Object Type", "Count"] + RATINGS + ["Total Lines", "Avg Complexity"]
    for j, h in enumerate(hdr, 1):
        c = put(ws, r0, j, h, HDR_FONT, Alignment(horizontal="center", vertical="center", wrap_text=True), NAVY)
        c.border = BORDER
    shown = [t for t in TYPE_ORDER if t != "External Table" or any(o["object_type"] == t for o in objects)]
    for i, t in enumerate(shown, r0 + 1):
        put(ws, i, 1, t, Font(name=FONT, bold=True, size=10))
        put(ws, i, 2, f'=COUNTIF(Inventory!$E:$E,$A{i})', align=Alignment(horizontal="center"))
        for j, rt in enumerate(RATINGS, 3):
            put(ws, i, j, f'=COUNTIFS(Inventory!$E:$E,$A{i},Inventory!$H:$H,"{rt}")', align=Alignment(horizontal="center"))
        put(ws, i, 7, f'=SUMIF(Inventory!$E:$E,$A{i},Inventory!$F:$F)', align=Alignment(horizontal="center"), fmt="#,##0")
        put(ws, i, 8, f'=IF(B{i}=0,"-",AVERAGEIF(Inventory!$E:$E,$A{i},Inventory!$G:$G))',
            align=Alignment(horizontal="center"), fmt="0.0")
        for j in range(1, 9):
            ws.cell(row=i, column=j).border = BORDER
            if (i - r0) % 2 == 0:
                ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=LIGHT)
    tot = r0 + len(shown) + 1
    put(ws, tot, 1, "Total", Font(name=FONT, bold=True, size=10, color="FFFFFF"), fill=TEAL)
    for j in range(2, 8):
        L = get_column_letter(j)
        put(ws, tot, j, f"=SUM({L}{r0 + 1}:{L}{tot - 1})", Font(name=FONT, bold=True, size=10, color="FFFFFF"),
            Alignment(horizontal="center"), TEAL, fmt="#,##0")
    put(ws, tot, 8, f'=IF(B{tot}=0,"-",AVERAGE(Inventory!$G:$G))', Font(name=FONT, bold=True, size=10, color="FFFFFF"),
        Alignment(horizontal="center"), TEAL, fmt="0.0")
    for j, rt in enumerate(RATINGS, 3):
        ws.cell(row=r0, column=j).fill = PatternFill("solid", fgColor=RATING_FILL[rt])
        ws.cell(row=r0, column=j).font = Font(name=FONT, bold=True, size=10, color=RATING_FONT[rt])
    for j in range(1, 9):
        ws.cell(row=tot, column=j).border = BORDER

    # chart
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.title = "Migration complexity by object type"
    ch.y_axis.title = "Objects"
    ch.height, ch.width = 7.5, 16
    ch.x_axis.scaling.orientation = "maxMin"   # same order as the table
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines = None
    data = Reference(ws, min_col=3, max_col=6, min_row=r0, max_row=tot - 1)
    cats = Reference(ws, min_col=1, min_row=r0 + 1, max_row=tot - 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    for s, rt in zip(ch.series, RATINGS):
        s.graphicalProperties.solidFill = {"Low": "70AD47", "Medium": "FFC000", "High": "ED7D31", "Very High": "C00000"}[rt]
        s.graphicalProperties.line.noFill = True
    ws.add_chart(ch, f"A{tot + 2}")

    # feature usage
    fr = r0
    put(ws, fr - 1, 10, "T-SQL patterns needing migration attention", Font(name=FONT, bold=True, size=12, color=NAVY))
    for j, h in enumerate(["Pattern", "Objects", "% of Objects"], 10):
        c = put(ws, fr, j, h, HDR_FONT, Alignment(horizontal="center", wrap_text=True), NAVY)
        c.border = BORDER
    ordered = sorted(feature_keys, key=lambda k: -sum(k in o.get("features", []) for o in objects))
    for i, k in enumerate(ordered, fr + 1):
        col = get_column_letter(4 + feature_keys.index(k))
        put(ws, i, 10, FEATURE_BY_KEY[k][1])
        put(ws, i, 11, f"=COUNTIF('Feature Matrix'!{col}:{col},\"Y\")", align=Alignment(horizontal="center"))
        put(ws, i, 12, f"=IF(COUNTA(Inventory!$A:$A)<=1,0,K{i}/(COUNTA(Inventory!$A:$A)-1))",
            align=Alignment(horizontal="center"), fmt="0%")
        for j in (10, 11, 12):
            ws.cell(row=i, column=j).border = BORDER
    from openpyxl.formatting.rule import DataBarRule
    ws.conditional_formatting.add(f"K{fr + 1}:K{fr + len(ordered)}",
                                  DataBarRule(start_type="num", start_value=0, end_type="max", color="5B9BD5"))

    # top complex objects (snapshot)
    tr = max(tot + 18, fr + len(ordered) + 3)
    put(ws, tr - 1, 1, "Top 15 most complex objects", Font(name=FONT, bold=True, size=12, color=NAVY))
    for j, h in enumerate(["Object", "Type", "", "Lines", "Score", "Rating"], 1):
        c = put(ws, tr, j, h or None, HDR_FONT, Alignment(horizontal="center"), NAVY)
        c.border = BORDER
    ws.merge_cells(start_row=tr, start_column=2, end_row=tr, end_column=3)
    top = sorted(objects, key=lambda o: -o["complexity_score"])[:15]
    for i, o in enumerate(top, tr + 1):
        put(ws, i, 1, o["full_name"])
        put(ws, i, 2, o["object_type"])
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
        put(ws, i, 4, o["lines"], align=Alignment(horizontal="center"))
        put(ws, i, 5, o["complexity_score"], align=Alignment(horizontal="center"), fmt="0.0")
        put(ws, i, 6, o["rating"], align=Alignment(horizontal="center"))
        for j in range(1, 7):
            ws.cell(row=i, column=j).border = BORDER
    rating_cf(ws, "F", tr + 1, tr + len(top))


def write_about(wb, meta):
    ws = wb.create_sheet("About")
    ws.sheet_properties.tabColor = "D9D9D9"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    put(ws, 1, 1, "How this report was produced", Font(name=FONT, bold=True, size=14, color=NAVY))
    rows = [
        ("Object types", "Stored procedures, functions (scalar / inline TVF / multi-statement TVF), views, materialized "
                         "views, and Databricks mirror objects (external tables or views reading Delta data produced by Databricks)."),
        ("Databricks mirror rule", f"An external table or view is classed as a Databricks mirror when its file format is DELTA, "
                                   f"it uses OPENROWSET(... FORMAT='DELTA'), or its name / data source / location matches the "
                                   f"pattern: {meta.get('mirror_pattern')}. Check the 'Mirror Evidence' column."),
        ("Complexity score", "min(non-blank lines / 25, 20) + sum of pattern weights + 0.5 per referenced object (max 8)."),
        ("Rating bands", " · ".join(f"{lbl} ≤ {lim}" if lim < 10 ** 8 else f"{lbl} > {RATING_BANDS[-2][0]}" for lim, lbl in RATING_BANDS)),
        ("Dependencies", "Parsed from the T-SQL text (FROM / JOIN / INSERT / UPDATE / MERGE / DELETE / EXEC / function calls); "
                         "in live mode merged with sys.sql_expression_dependencies. 'Resolved' = the referenced object is "
                         "itself in this inventory; unresolved names are usually tables."),
        ("Limitations", "Regex-based static analysis: names built in dynamic SQL are not detected, and comments/strings are "
                        "ignored. Treat ratings as a triage aid for estimation, not a final effort figure."),
    ]
    r = 3
    for k, v in rows:
        put(ws, r, 1, k, Font(name=FONT, bold=True, size=10)); put(ws, r, 2, v, BODY, WRAP_TOP); r += 1
    r += 1
    put(ws, r, 1, "Pattern weights", Font(name=FONT, bold=True, size=12, color=NAVY)); r += 1
    for j, h in enumerate(["Pattern", "Weight"], 1):
        put(ws, r, j, h, HDR_FONT, fill=NAVY)
    r += 1
    for k, label, _, w, *_ in FEATURES:
        put(ws, r, 1, label); put(ws, r, 2, w, align=Alignment(horizontal="left")); r += 1
    if meta.get("warnings"):
        r += 1
        put(ws, r, 1, "Extraction warnings", Font(name=FONT, bold=True, size=12, color="C00000")); r += 1
        for w in meta["warnings"]:
            put(ws, r, 2, w, BODY, WRAP_TOP); r += 1


# --------------------------------------------------------------------------- lineage
def write_lineage(wb, objects, nodes, edges, waves, cyclic, anchors, html_name, images, depth, max_diagrams):
    import math
    import tempfile
    from openpyxl.drawing.image import Image as XLImage

    out_, in_ = lin._adj(edges)
    inv = [o for o in objects if o["full_name"] in nodes]

    # ---- Lineage (table)
    ws = wb.create_sheet("Lineage")
    ws.sheet_properties.tabColor = "1F3864"
    title_block(ws, "Object lineage", "Upstream = objects/tables that feed this object · Downstream = everything "
                "impacted if it changes · Migration wave = suggested build/load order (wave 1 first).")
    if html_name:
        c = put(ws, 3, 1, f"Open the interactive lineage explorer: {html_name}  (same folder as this workbook)",
                Font(name=FONT, size=10, color="0563C1", underline="single"))
        c.hyperlink = html_name
    if not images:
        put(ws, 4, 1, "Diagram images were not added to this workbook (matplotlib not available or turned off) - "
            "use the interactive lineage explorer above for the diagrams.",
            Font(name=FONT, italic=True, size=9, color="595959"))
    headers = ["#", "Object (→ diagram)", "Type", "Rating", "Migration Wave", "Direct Upstream", "Direct Downstream",
               "Upstream (all)", "Downstream (all)", "Circular Dependency", "Definition"]
    header_row(ws, 5, headers, [5, 40, 18, 11, 10, 44, 44, 10, 11, 11, 12])
    diag_rows = {}
    r = 6
    for i, o in enumerate(inv, 1):
        fn = o["full_name"]
        put(ws, r, 1, i, BODY, TOP)
        link_c = put(ws, r, 2, fn, Font(name=FONT, size=10, color="0563C1", underline="single"), TOP)
        diag_rows[fn] = link_c
        put(ws, r, 3, o["object_type"], BODY, TOP)
        put(ws, r, 4, o.get("rating"), BODY, TOP)
        put(ws, r, 5, waves.get(fn), BODY, Alignment(horizontal="center", vertical="top"))
        put(ws, r, 6, txt(sorted(set(in_.get(fn, [])))), BODY, WRAP_TOP)
        put(ws, r, 7, txt(sorted(set(out_.get(fn, [])))), BODY, WRAP_TOP)
        put(ws, r, 8, len(lin.closure(fn, in_)), BODY, Alignment(horizontal="center", vertical="top"))
        put(ws, r, 9, len(lin.closure(fn, out_)), BODY, Alignment(horizontal="center", vertical="top"))
        put(ws, r, 10, "Y" if fn in cyclic else None, BODY, Alignment(horizontal="center", vertical="top"))
        d = put(ws, r, 11, "view code" if anchors.get(fn) else None,
                Font(name=FONT, size=10, color="0563C1", underline="single"), TOP)
        if anchors.get(fn):
            d.hyperlink = f"#Definitions!B{anchors[fn]}"
        r += 1
    add_table(ws, "tblLineage", 5, r - 1, len(headers))
    rating_cf(ws, "D", 6, r - 1)
    ws.conditional_formatting.add(f"J6:J{max(r - 1, 6)}", CellIsRule(operator="equal", formula=['"Y"'],
                                  fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.freeze_panes = "C6"

    # ---- Lineage Diagrams (one block per object)
    if not images:
        for c in diag_rows.values():
            c.font = BODY
        return
    wd = wb.create_sheet("Lineage Diagrams")
    wd.sheet_properties.tabColor = "1F3864"
    wd.sheet_view.showGridLines = False
    wd.column_dimensions["A"].width = 46
    wd.column_dimensions["B"].width = 2
    title_block(wd, "Lineage diagrams", f"Up to {depth} hops each way around the selected object (bold border). "
                "Arrows read 'is needed by': solid = data read/written, dashed = procedure called by, dotted = function used by.")
    tmp = tempfile.mkdtemp(prefix="lineage_png_")
    row, drawn = 4, 0
    for o in inv:
        fn = o["full_name"]
        if not (in_.get(fn) or out_.get(fn)):
            diag_rows[fn].font = BODY
            continue
        if drawn >= max_diagrams:
            diag_rows[fn].font = BODY
            continue
        sub_nodes, sub_edges = lin.neighbourhood(fn, nodes, edges, up=depth, down=depth)
        png = lin.render_png(fn, sub_nodes, sub_edges, os.path.join(tmp, f"{drawn:05d}.png"))
        if not png:
            wd.cell(row=row, column=1, value="matplotlib not installed - diagrams skipped (pip install matplotlib)")
            break
        drawn += 1
        for col in range(1, 16):
            wd.cell(row=row, column=col).fill = HDR_FILL
        put(wd, row, 1, f"{fn}   ·   {o['object_type']}   ·   wave {waves.get(fn)}"
            + (f"   ·   {o.get('rating')}" if o.get("rating") else ""),
            Font(name=FONT, bold=True, color="FFFFFF", size=10), fill=NAVY)
        diag_rows[fn].hyperlink = f"#'Lineage Diagrams'!A{row}"
        img = XLImage(png)
        max_w = 1100
        if img.width > max_w:
            img.height = int(img.height * max_w / img.width)
            img.width = max_w
        wd.add_image(img, f"C{row + 1}")
        ups = sorted(set(in_.get(fn, [])))
        downs = sorted(set(out_.get(fn, [])))
        text = (f"Upstream ({len(ups)})\n" + ("\n".join("  ← " + u for u in ups) or "  none")
                + f"\n\nDownstream ({len(downs)})\n" + ("\n".join("  → " + d for d in downs) or "  none"))
        put(wd, row + 1, 1, txt(text), BODY, WRAP_TOP)
        text_lines = text.count("\n") + 1
        img_rows = math.ceil(img.height / 20) + 1
        span = max(img_rows, math.ceil(text_lines * 15 / 20) + 1, 4)
        wd.merge_cells(start_row=row + 1, start_column=1, end_row=row + span, end_column=1)
        row += span + 2


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="inp", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--target", choices=["databricks", "fabric"], default="databricks")
    ap.add_argument("--client")
    ap.add_argument("--no-definitions", action="store_true")
    ap.add_argument("--no-lineage", action="store_true", help="skip Lineage sheets and the HTML explorer")
    ap.add_argument("--no-lineage-images", action="store_true", help="Lineage table only, no diagram images")
    ap.add_argument("--lineage-depth", type=int, default=2, help="hops each way in Excel diagrams (default 2)")
    ap.add_argument("--max-diagrams", type=int, default=400, help="cap on diagram images in the workbook")
    ap.add_argument("--lineage-html", help="path for the interactive explorer (default: <out>_Lineage.html)")
    a = ap.parse_args()

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    data = json.load(open(a.inp, encoding="utf-8"))
    meta, objects = data["meta"], data["objects"]
    objects.sort(key=lambda o: (TYPE_ORDER.index(o["object_type"]) if o["object_type"] in TYPE_ORDER else 99,
                                o["schema"].lower(), o["name"].lower()))
    deps = analyze(objects, a.target)

    wb = Workbook()
    anchors = {} if a.no_definitions else None
    feature_keys = [f[0] for f in FEATURES]
    write_summary(wb, meta, objects, feature_keys, a.target, a.client)
    if anchors is None:
        anchors = write_definitions(wb, objects)
    write_inventory(wb, objects, anchors)
    for t in TYPE_ORDER:
        objs = [o for o in objects if o["object_type"] == t]
        if objs or t != "External Table":
            write_type_sheet(wb, t, objs, anchors)
    write_features(wb, objects)
    write_dependencies(wb, deps)
    html_path, notes = None, []
    if not a.no_lineage:
        nodes, edges = lin.build_graph(objects, deps)
        waves, cyclic = lin.migration_waves(nodes, edges)
        for o in objects:
            o["migration_wave"] = waves.get(o["full_name"])
        html_path = a.lineage_html or os.path.splitext(a.out)[0] + "_Lineage.html"
        lin.write_html(nodes, edges, waves, cyclic, meta, html_path)
        try:
            rel = os.path.relpath(html_path, os.path.dirname(os.path.abspath(a.out)) or ".")
        except ValueError:   # different drive on Windows
            rel = os.path.abspath(html_path)
        images = not a.no_lineage_images
        if images and not _bootstrap.has_matplotlib():
            images = False
            notes.append("matplotlib not available: Lineage Diagrams sheet skipped - the HTML explorer has all diagrams")
        write_lineage(wb, objects, nodes, edges, waves, cyclic, anchors, rel, images,
                      a.lineage_depth, a.max_diagrams)
    write_about(wb, meta)
    # order: Summary, Inventory, type sheets, Feature Matrix, Dependencies, Definitions, About
    if "Definitions" in wb.sheetnames:
        d = wb["Definitions"]
        wb.move_sheet(d, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Definitions") - 1)
    wb.calculation.fullCalcOnLoad = True   # Excel computes formulas on open; no LibreOffice needed
    wb.save(a.out)
    counts = {}
    for o in objects:
        counts[o["object_type"]] = counts.get(o["object_type"], 0) + 1
    print(json.dumps({"out": a.out, "objects": len(objects), "by_type": counts, "dependencies": len(deps),
                      "lineage_html": html_path, "notes": notes}, indent=1, ensure_ascii=False))


if __name__ == "__main__":
    main()
